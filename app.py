import streamlit as st
import pandas as pd
from docx import Document
from docx.shared import Pt
import io


def get_checklist_structure():
    """
    Returns the checklist data in a structured, nested dictionary.
    This makes rendering and processing the data much cleaner.
    """
    return {
        "1. TORAX PA y LAT": {
            "Criterios de imagen": [
                "Visualización detallada del patrón vascular en todo el pulmón, especialmente los vasos periféricos.",
                "Visualización detallada de la tráquea y los bronquios principales.",
                "Visualización detallada de los contornos cardiaco y aórtico.",
                "Visualización detallada de los diafragmas y los ángulos costofrénicos laterales.",
                "Visualización del parénquima pulmonar retrocardíaco, mediastino, y columna a través del corazón.",
            ],
            "Detalles importantes de la imagen": [
                "Nódulos (Alto contraste: 0,7 mm de diámetro).",
                "Nódulos (Bajo contraste: 2 mm de diámetro).",
                "Detalles lineales y reticulares (Alto contraste: 0,3 mm de espesor).",
                "Detalles lineales y reticulares (Bajo contraste: 2 mm de espesor).",
            ],
        },
        "2. CRÁNEO PA y LATERAL": {
            "Criterios de imagen": [
                "Visualización de senos frontales, celdas etmoidales, punta del peñasco y conductos auditivos internos.",
                "Buena definición de las tablas interna y externa de la bóveda craneal.",
                "Reproducción nítida de surcos vasculares, vértex y estructura trabecular del cráneo (lateral).",
            ],
            "Detalles importantes de la imagen": [
                "Detalles (Cráneo): 0,3-0,5 mm",
            ],
        },
        "3. COLUMNA LUMBAR AP/PA y LATERAL": {
            "Criterios de imagen": [
                "Reproducción nítida de los platillos superior e inferior de la vértebra central.",
                "Una única imagen en el borde posterior vertebral (lateral).",
                "Correcta visualización de los pedículos (AP o PA).",
                "Visualización de las articulaciones intervertebrales.",
                "Visualización de la apófisis espinosas y transversas.",
                "Reproducción nítida de la cortical y la estructura trabecular.",
                "Visualización de las partes blandas adyacentes (contornos del psoas).",
                "Reproducción de las articulaciones sacroilíacas (AP o PA).",
            ],
            "Detalles importantes de la imagen": [
                "Detalles (Columna Lumbar): 0,3-0,5 mm",
            ],
        },
        "4. PELVIS": {
            "Criterios de imagen": [
                "Visualización detallada del sacro y los agujeros de conjunción.",
                "Reproducción correcta de la cortical y la esponjosa de los trocánteres.",
            ],
            "Detalles importantes de la imagen": [
                "Detalles: 0,5 mm",
            ],
        },
        "5. ABDOMEN AP Y APARATO URINARIO": {
            "Criterios de imagen": [
                "Visualización de los contornos renales.",
                "Visualización de los contornos del psoas.",
                "Correcta reproducción de los huesos.",
                "Aumento de densidad del parénquima (efecto nefrográfico).",
                "Correcta visualización de la pelvis renal y de los cálices.",
                "Visualización de la unión pielo-ureteral y de todo el trayecto de los uréteres.",
                "Reproducción de toda el área vesical.",
                "Borde hepático",
                "Borde esplénico"
            ],
            "Detalles importantes de la imagen": [
                "Detalles caliciales: 0,3 mm.",
                "Calcificaciones: 1 mm",
            ],
        },
    }


def create_word_summary(structure, states, filename):
    """
    Generates a .docx summary of the checklist.

    Args:
        filename (str or file-like object): The path or buffer to save the document to.
    """
    doc = Document()
    doc.add_heading("Resumen de Checklist de Calidad de Imagen", level=0)

    for header, sub_sections in structure.items():
        doc.add_heading(header, level=1)
        for sub_header, items_list in sub_sections.items():
            doc.add_heading(sub_header, level=2)
            for item in items_list:
                status_symbol = "✓" if states.get(item, False) else "✗"
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{status_symbol} ").bold = True
                p.add_run(item)
    doc.save(filename)


def create_app():
    """
    Sets up and runs the Streamlit checklist application.
    """
    st.set_page_config(page_title="Validación de Calidad de Imagen", layout="centered")

    st.title("Checklist de Validación de la Calidad de Imagen Clínica")
    st.write(
        "Clica las casillas para marcar los criterios de calidad de imagen clínica que has podido visualizar.",
        "Cuando hayas terminado, pulsa el botón de abajo para generar los archivos de resultados.",
        "Podrás descargar un archivo Excel (.xlsx) con los datos y un resumen en formato Word (.docx)."
    )

    # --- Checklist Data ---
    checklist_structure = get_checklist_structure()

    # Use session_state to preserve checkbox states across reruns
    if 'checkbox_states' not in st.session_state:
        # Flatten the structure to initialize the state dictionary
        all_items = [
            item
            for section in checklist_structure.values()
            for sub_section in section.values()
            for item in sub_section
        ]
        st.session_state.checkbox_states = {item: False for item in all_items}

    # Initialize session state for file generation logic
    if "files_generated" not in st.session_state:
        st.session_state.files_generated = False
        st.session_state.excel_buffer = None
        st.session_state.word_buffer = None

    # --- Render sections with checkboxes ---
    for header, sub_sections in checklist_structure.items():
        st.header(header)
        for sub_header, items_list in sub_sections.items():
            st.subheader(sub_header)
            for item in items_list:
                st.session_state.checkbox_states[item] = st.checkbox(
                    item, value=st.session_state.checkbox_states.get(item, False), key=item
                )

    # --- Save and Generate Files ---
    if st.button("Guardar y Generar Archivos"):
        data_to_save = {
            "Criterio": list(st.session_state.checkbox_states.keys()),
            "Completado": list(st.session_state.checkbox_states.values())
        }
        df = pd.DataFrame(data_to_save)

        # --- Generate Files in Memory and store in session state ---
        excel_buffer = io.BytesIO()
        df.to_excel(excel_buffer, index=False, engine="openpyxl")
        st.session_state.excel_buffer = excel_buffer

        word_buffer = io.BytesIO()
        create_word_summary(checklist_structure, st.session_state.checkbox_states, word_buffer)
        st.session_state.word_buffer = word_buffer

        st.session_state.files_generated = True

    # --- Download Buttons ---
    # Always show download buttons if files have been generated
    if st.session_state.files_generated:
        st.success("¡Archivos generados con éxito!")

        col1, col2 = st.columns(2)
        with col1:
            st.download_button(
                label="Descargar Excel (.xlsx)",
                data=st.session_state.excel_buffer,
                file_name="checklist_status.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        with col2:
            st.download_button(
                label="Descargar Resumen (.docx)",
                data=st.session_state.word_buffer,
                file_name="checklist_summary.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

if __name__ == "__main__":
    create_app()
